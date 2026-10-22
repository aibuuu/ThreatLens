import io
import json
import re
from datetime import datetime
from html.parser import HTMLParser

from flask import Flask, render_template, request
from werkzeug.utils import secure_filename

try:
    import pdfplumber
except ImportError:
    pdfplumber = None
try:
    from pptx import Presentation
except ImportError:
    Presentation = None
try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 5 * 1024 * 1024
ALLOWED = (".txt", ".html", ".htm", ".json", ".pdf", ".pptx", ".xlsx")
APPROVED_STORE = []  # stand-in for the vector DB; swap for ChromaDB later


def gap(n):
    """Match up to n characters of anything, including newlines."""
    return r"[\s\S]{0,%d}?" % n


# ---------------------------------------------------------------------------
# Rules
# kind: "threat" (counts toward quarantine), "link", "note" (low-level instruction)
# All scan rules are case-insensitive and newline-tolerant.
# ---------------------------------------------------------------------------
SCAN_RULES = [
    ("Instruction Override", "HIGH", 30, "threat",
     r"\b(?:ignore|disregard|forget|override|bypass)\s+(?:(?:all|any|the|your|these|those|of)\s+)*"
     r"(?:(?:previous|prior|earlier|above|preceding|original|existing|system|safety|conflicting)\s+)?"
     r"(?:instructions?|rules|guidelines?|guidance|directions?|directives?|prompts?)\b"),
    ("System Override", "HIGH", 30, "threat",
     r"\bsystem\s+override\b|\boverride\s+(?:the\s+)?(?:system|safety)\b"),
    ("New Instruction Header", "HIGH", 25, "threat",
     r"\bnew\s+instructions?\s*:"),
    ("Role Manipulation", "MEDIUM", 15, "threat",
     r"\b(?:act\s+as|pretend\s+to\s+be|you\s+are\s+now|roleplay\s+as)\s+(?:an?\s+|the\s+)?"
     r"(?:admin|administrator|developer|system|root|unrestricted|unfiltered)\b"),
    ("System Prompt Disclosure", "HIGH", 30, "threat",
     r"\b(?:reveal|show|print|expose|display|repeat|disclose|leak|output)\b" + gap(40) +
     r"\b(?:system\s+prompt|hidden\s+instructions?|internal\s+prompt|initial\s+instructions?|your\s+setup)\b"),
    ("Credential Disclosure", "HIGH", 30, "threat",
     r"\b(?:reveal|disclose|expose|print|dump|output|share)\b" + gap(40) +
     r"\b(?:passwords?|secrets?|api\s+keys?|access\s+tokens?|credentials)\b"),
    ("Security Bypass", "HIGH", 25, "threat",
     r"\b(?:bypass|disable|evade|circumvent|turn\s+off|deactivate)\b" + gap(40) +
     r"\b(?:security|safety|authentication|authorization|access\s+controls?|safeguards?|guardrails?|safety\s+filters?)\b"),
    ("Confidential Data Transfer", "HIGH", 25, "threat",
     r"\b(?:send|share|transfer|upload|forward|export|transmit)\b" + gap(60) +
     r"\b(?:confidential|private|sensitive|personal)\b" + gap(30) +
     r"\b(?:data|information|files?|records?)\b"),
    ("Outbound Email Exfiltration", "MEDIUM", 15, "threat",
     r"\b(?:forward|send|upload|transmit)\b" + gap(80) + r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+"),
    ("Data Exfiltration", "MEDIUM", 15, "threat",
     r"\b(?:extract|dump|exfiltrate|harvest|copy|download)\b" + gap(50) +
     r"\b(?:data|records?|credentials|passwords|tokens)\b"),
    ("Instruction Concealment", "MEDIUM", 15, "threat",
     r"\b(?:do\s+not|don'?t|never)\s+(?:tell|mention|inform|reveal)\b" + gap(50) +
     r"\b(?:the\s+user|the\s+reviewer|the\s+owner|this\s+instruction|these\s+instructions|this\s+note)\b"),
    ("Audit Log Deletion", "HIGH", 25, "threat",
     r"\b(?:delete|erase|clear|remove)\b" + gap(30) +
     r"\b(?:audit\s+logs?|security\s+logs?|activity\s+logs?)\b"),
    ("Indirect Prompt Injection", "HIGH", 25, "threat",
     r"\b(?:when\s+you\s+read\s+this|if\s+an?\s+ai\s+reads\s+this|assistants?|ai\s+models?|language\s+models?)\b"
     + gap(40) + r"\b(?:must|should)\s+(?:ignore|reveal|disclose|obey|execute)\b"),
    ("Summarizer Instruction", "LOW", 10, "note",
     r"\b(?:note|update|message|instruction|directive)s?\s+(?:to|for)\s+(?:ai\s+|automated\s+)?"
     r"(?:assistants?|summarizers?|ai\s+models?)\b"),
    ("Assistant Directive Marker", "LOW", 10, "note",
     r"\bassistant\s+directive\b|\bsystem\s+note\b"),
    # Links
    ("Phishing-style Link", "HIGH", 25, "link",
     r"https?://[\w.-]*(?:login|verify|secure|update|reset)[\w.-]*\.(?:xyz|top|tk|ml|click|zip)\b"),
    ("Raw IP Address URL", "MEDIUM", 15, "link",
     r"https?://\d{1,3}(?:\.\d{1,3}){3}[^\s\"'<>)]*"),
    ("URL Shortener", "LOW", 5, "link",
     r"\b(?:bit\.ly|tinyurl\.com|goo\.gl)/[^\s\"'<>)]+"),
]


def luhn_ok(value):
    digits = [int(c) for c in re.sub(r"\D", "", value)]
    if not 13 <= len(digits) <= 19:
        return False
    total = 0
    for i, d in enumerate(reversed(digits)):
        if i % 2:
            d *= 2
            d -= 9 if d > 9 else 0
        total += d
    return total % 10 == 0


def always(_value):
    return True


# (kind, name, severity, points, pattern, flags, redaction token, validator)
DATA_RULES = [
    ("pii", "Card Number", "HIGH", 20, r"(?<!\d)(?:\d[ -]?){12,18}\d(?!\d)", 0, "[CARD_REDACTED]", luhn_ok),
    ("pii", "Aadhaar Number", "HIGH", 20, r"(?<!\d)[2-9]\d{3}[ -]?\d{4}[ -]?\d{4}(?!\d)", 0, "[AADHAAR_REDACTED]", always),
    ("pii", "PAN", "MEDIUM", 15, r"\b[A-Z]{5}[0-9]{4}[A-Z]\b", 0, "[PAN_REDACTED]", always),
    ("pii", "IFSC Code", "MEDIUM", 15, r"\b[A-Z]{4}0[A-Z0-9]{6}\b", 0, "[IFSC_REDACTED]", always),
    ("pii", "Bank Account Number", "MEDIUM", 15,
     r"\baccounts?(?:\s+(?:no\.?|number))?\s*[:#]?\s*\d{9,18}\b", re.I, "[ACCOUNT_REDACTED]", always),
    ("pii", "Indian Mobile Number", "LOW", 5,
     r"(?<!\d)(?:\+?91[-\s]?)?[6-9]\d{4}[-\s]?\d{5}(?!\d)", 0, "[PHONE_REDACTED]", always),
    ("pii", "Short Phone Number", "LOW", 5, r"(?<!\d)\d{3}-\d{4}(?!\d)", 0, "[PHONE_REDACTED]", always),
    ("pii", "Email Address", "LOW", 5, r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+", 0, "[EMAIL_REDACTED]", always),
    ("secret", "AWS Access Key", "HIGH", 30, r"\bAKIA[0-9A-Z]{16}\b", 0, "[API_KEY_REDACTED]", always),
    ("secret", "API Key Token", "HIGH", 25,
     r"\b(?:sk-[A-Za-z0-9_\-]{16,}|ghp_[A-Za-z0-9]{20,})\b", 0, "[API_KEY_REDACTED]", always),
    ("secret", "Secret Assignment", "HIGH", 25,
     r"\b[a-z_]*(?:secret|token|api[_-]?key|password)[a-z_]*\s*[=:]\s*\S{12,}", re.I, "[SECRET_REDACTED]", always),
]

HIDDEN_STYLE = re.compile(
    r"display\s*:\s*none|visibility\s*:\s*hidden|color\s*:\s*(?:#fff(?:fff)?|white)\b|font-size\s*:\s*[0-5]px",
    re.I)
VOID_TAGS = {"meta", "br", "img", "input", "link", "hr", "source", "wbr"}

ACTIONS = {
    "QUARANTINE": ("Quarantine: do not ingest",
                   "Hold this document for human review. It will not be stored in the approved set."),
    "SANITIZE": ("Sanitize before ingesting",
                 "A redacted copy is produced: suspicious sentences are removed and sensitive values replaced. "
                 "Review it before storing."),
    "ALLOW": ("Allow",
              "No rule matched. This is not a safety guarantee."),
}
SUMMARIES = {
    "QUARANTINE": "High-priority threat patterns found. This document should not be ingested.",
    "SANITIZE": "Sensitive or suspicious content found. A sanitized copy is available.",
    "ALLOW": "No detection rule matched this document.",
}


# ---------------------------------------------------------------------------
# Scanning
# ---------------------------------------------------------------------------
def quoted(text, pos):
    line_start = text.rfind("\n", 0, pos) + 1
    return text[line_start:pos].count('"') % 2 == 1


def context_of(text, start, end):
    line_start = text.rfind("\n", 0, start) + 1
    line_end = text.find("\n", end)
    if line_end == -1:
        line_end = len(text)
    return text[line_start:line_end].strip()[:300]


def mask(value):
    if len(value) <= 6:
        return "***"
    return value[:2] + "*" * (len(value) - 4) + value[-2:]


def make_finding(name, severity, points, kind, location, text, m):
    if m is None:
        matched, line_label, context = text.strip()[:120], location, text.strip()[:300]
    else:
        matched = m.group(0)
        line_no = text.count("\n", 0, m.start()) + 1
        line_label = f"{location} · L{line_no}"
        context = context_of(text, *m.span())
        if kind in ("pii", "secret"):
            masked = mask(matched)
            matched = masked
            context = context.replace(m.group(0), masked)
    return {"name": name, "severity": severity, "points": points, "kind": kind,
            "line": line_label, "matched_text": matched, "context": context}


def scan_unit(location, text, hidden):
    findings, threat_spans, data_spans = [], [], []

    def overlaps(span, spans):
        return any(span[0] < e and s < span[1] for s, e in spans)

    for name, severity, points, kind, pattern in SCAN_RULES:
        for m in re.finditer(pattern, text, re.I):
            span = m.span()
            if overlaps(span, threat_spans):
                continue
            threat_spans.append(span)
            if kind == "threat" and quoted(text, span[0]):
                # Quoted examples (e.g. security training) are shown but do not count.
                findings.append(make_finding(f"{name} (quoted, not counted)", severity, 0,
                                             "quoted", location, text, m))
                continue
            findings.append(make_finding(name, severity, points, kind, location, text, m))

    for kind, name, severity, points, pattern, flags, _token, valid in DATA_RULES:
        for m in re.finditer(pattern, text, flags):
            span = m.span()
            if not valid(m.group(0)) or overlaps(span, data_spans):
                continue
            data_spans.append(span)
            findings.append(make_finding(name, severity, points, kind, location, text, m))

    if hidden and text.strip():
        findings.append(make_finding("Hidden Content", "MEDIUM", 15, "hidden", location, text, None))
    return findings


def sanitize_text(text):
    kept = []
    for part in re.split(r"(?<=[.!?])\s+|\n+", text):
        part = part.strip()
        if not part:
            continue
        if any(re.search(p, part, re.I) for _n, _s, _pts, kind, p in SCAN_RULES
               if kind in ("threat", "link", "note")):
            kept.append("[REMOVED: suspicious instruction]")
        else:
            kept.append(part)
    out = " ".join(kept)
    for _kind, _name, _sev, _pts, pattern, flags, token, valid in DATA_RULES:
        out = re.sub(pattern, lambda m, t=token, v=valid: t if v(m.group(0)) else m.group(0),
                     out, flags=flags)
    return out


# ---------------------------------------------------------------------------
# Extraction: each function returns a list of (location, text, hidden) units
# ---------------------------------------------------------------------------
def is_white(color):
    return (isinstance(color, (list, tuple)) and len(color) in (1, 3)
            and all(isinstance(c, (int, float)) and c >= 0.95 for c in color))


def hidden_pdf_char(obj):
    return obj.get("object_type") == "char" and (
        is_white(obj.get("non_stroking_color")) or (obj.get("size") or 10) < 4)


def pdf_units(data):
    if pdfplumber is None:
        raise RuntimeError("pdfplumber is not installed (pip install pdfplumber)")
    units = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        meta = pdf.metadata or {}
        for key in ("Title", "Author", "Subject", "Keywords"):
            if meta.get(key):
                units.append((f"PDF metadata ({key})", str(meta[key]), False))
        for n, page in enumerate(pdf.pages, 1):
            visible = page.filter(lambda o: not hidden_pdf_char(o)).extract_text() or ""
            hidden = page.filter(hidden_pdf_char).extract_text() or ""
            units.append((f"PDF page {n}", visible, False))
            if hidden.strip():
                units.append((f"PDF page {n} (hidden text)", hidden, True))
    return units


def run_is_hidden(run):
    if run.font.size is not None and run.font.size.pt < 6:
        return True
    try:
        return str(run.font.color.rgb) == "FFFFFF"
    except Exception:
        return False


def pptx_units(data):
    if Presentation is None:
        raise RuntimeError("python-pptx is not installed (pip install python-pptx)")
    prs = Presentation(io.BytesIO(data))
    units = []
    for i, slide in enumerate(prs.slides, 1):
        visible, hidden = [], []
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                vis = "".join(r.text for r in para.runs if not run_is_hidden(r))
                hid = "".join(r.text for r in para.runs if run_is_hidden(r))
                if vis.strip():
                    visible.append(vis)
                if hid.strip():
                    hidden.append(hid)
        units.append((f"Slide {i}", "\n".join(visible), False))
        if hidden:
            units.append((f"Slide {i} (hidden text)", "\n".join(hidden), True))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text
            if notes.strip():
                units.append((f"Slide {i} speaker notes", notes, False))
    return units


def xlsx_units(data):
    if load_workbook is None:
        raise RuntimeError("openpyxl is not installed (pip install openpyxl)")
    wb = load_workbook(io.BytesIO(data), data_only=False)
    units = []
    for ws in wb.worksheets:
        sheet_hidden = ws.sheet_state != "visible"
        vis, hid = [], []
        for row in ws.iter_rows():
            cells = [str(c.value) for c in row if c.value is not None]
            if not cells:
                continue
            row_hidden = sheet_hidden or bool(ws.row_dimensions[row[0].row].hidden)
            (hid if row_hidden else vis).append(" | ".join(cells))
        if vis:
            units.append((f"Sheet '{ws.title}'", "\n".join(vis), False))
        if hid:
            units.append((f"Sheet '{ws.title}' (hidden)", "\n".join(hid), True))
    return units


class _HTMLSplit(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.stack = []           # (tag, hidden_here)
        self.css = []
        self.in_style = False
        self.in_script = False
        self.visible, self.hidden, self.comments = [], [], []

    def _hidden_classes(self):
        css = " ".join(self.css)
        return {c for c, rules in re.findall(r"\.([\w-]+)\s*\{([^}]*)\}", css) if HIDDEN_STYLE.search(rules)}

    def _is_hidden(self, attrs):
        d = dict(attrs)
        if HIDDEN_STYLE.search(d.get("style") or ""):
            return True
        return any(c in self._hidden_classes() for c in (d.get("class") or "").split())

    def handle_starttag(self, tag, attrs):
        if tag in VOID_TAGS:
            return
        if tag == "style":
            self.in_style = True
        elif tag == "script":
            self.in_script = True
        self.stack.append((tag, self._is_hidden(attrs)))

    def handle_endtag(self, tag):
        if tag == "style":
            self.in_style = False
        elif tag == "script":
            self.in_script = False
        for i in range(len(self.stack) - 1, -1, -1):
            if self.stack[i][0] == tag:
                del self.stack[i:]
                break

    def handle_data(self, data):
        if self.in_style:
            self.css.append(data)
        elif self.in_script:
            return
        elif any(h for _t, h in self.stack):
            self.hidden.append(data)
        else:
            self.visible.append(data)

    def handle_comment(self, data):
        if data.strip():
            self.comments.append(data.strip())


def html_units(text):
    p = _HTMLSplit()
    p.feed(text)
    p.close()
    units = []
    vis = re.sub(r"\s+", " ", " ".join(p.visible)).strip()
    if vis:
        units.append(("HTML body", vis, False))
    hid = re.sub(r"\s+", " ", " ".join(p.hidden)).strip()
    if hid:
        units.append(("HTML hidden element", hid, True))
    # Comments are scanned for injection but are not flagged as hidden by themselves.
    for c in p.comments:
        units.append(("HTML comment", c, False))
    return units


def json_units(text):
    obj = json.loads(text)
    units = []

    def walk(node, path, hidden):
        if isinstance(node, dict):
            for k, v in node.items():
                key = str(k)
                child_hidden = hidden or key.lower().startswith(("internal", "_", "hidden"))
                walk(v, f"{path}.{key}" if path else key, child_hidden)
        elif isinstance(node, list):
            for i, v in enumerate(node):
                walk(v, f"{path}[{i}]", hidden)
        elif isinstance(node, str):
            units.append((path or "value", node, hidden))

    walk(obj, "", False)
    return units, obj


def rebuild_json(node, path, parts):
    if isinstance(node, dict):
        return {k: rebuild_json(v, f"{path}.{k}" if path else str(k), parts) for k, v in node.items()}
    if isinstance(node, list):
        return [rebuild_json(v, f"{path}[{i}]", parts) for i, v in enumerate(node)]
    if isinstance(node, str):
        return parts.get(path or "value", node)
    return node


def build_units(lower, data):
    if lower.endswith(".pdf"):
        return pdf_units(data), None
    if lower.endswith(".pptx"):
        return pptx_units(data), None
    if lower.endswith(".xlsx"):
        return xlsx_units(data), None
    text = data.decode("utf-8-sig", errors="replace")
    if lower.endswith((".html", ".htm")):
        return html_units(text), None
    if lower.endswith(".json"):
        return json_units(text)
    return [("body", text, False)], None


# ---------------------------------------------------------------------------
# Verdict
# ---------------------------------------------------------------------------
def analyze(filename, data):
    lower = filename.lower()
    findings = []
    if lower.endswith(".pdf") and re.search(rb"/JavaScript|/JS\s*[(<]", data):
        findings.append(make_finding("Embedded JavaScript", "HIGH", 25, "threat",
                                     "document", "/JavaScript action", None))

    units, obj = build_units(lower, data)
    sanitized_parts = {}
    for location, text, hidden in units:
        findings += scan_unit(location, text, hidden)
        sanitized_parts[location] = sanitize_text(text)

    if obj is not None:
        sanitized = json.dumps(rebuild_json(obj, "", sanitized_parts), indent=2, ensure_ascii=False)
    elif len(units) == 1:
        sanitized = sanitized_parts[units[0][0]]
    else:
        sanitized = "\n\n".join(f"[{loc}]\n{sanitized_parts[loc]}" for loc, _t, _h in units)

    scored = [f for f in findings if f["points"] > 0]
    threat_points = sum(f["points"] for f in scored if f["kind"] in ("threat", "link"))
    score = min(100, sum(f["points"] for f in scored))

    if threat_points >= 25:
        verdict, risk = "QUARANTINE", "HIGH"
    elif scored:
        verdict, risk = "SANITIZE", "MEDIUM"
    else:
        verdict, risk = "ALLOW", "SAFE"

    return {
        "filename": filename,
        "verdict": verdict,
        "risk": risk,
        "score": score,
        "findings": sorted(findings, key=lambda f: f["points"], reverse=True),
        "high": sum(f["severity"] == "HIGH" for f in scored),
        "medium": sum(f["severity"] == "MEDIUM" for f in scored),
        "low": sum(f["severity"] == "LOW" for f in scored),
        "sanitized": "" if verdict == "QUARANTINE" else sanitized,
    }


# ---------------------------------------------------------------------------
# Web routes
# ---------------------------------------------------------------------------
def page_context(result, stored):
    title, action = ACTIONS[result["verdict"]]
    if stored:
        action += f" Stored in approved set ({len(APPROVED_STORE)} total)."
    return {
        "findings": result["findings"],
        "filename": result["filename"],
        "score": result["score"],
        "risk": result["risk"],
        "risk_class": result["risk"],
        "status": "SCANNED",
        "scanned_at": datetime.now().strftime("%d %b %Y · %H:%M"),
        "high_count": result["high"],
        "medium_count": result["medium"],
        "low_count": result["low"],
        "summary": SUMMARIES[result["verdict"]],
        "action_title": title,
        "action_text": action,
        "verdict": result["verdict"],
        "sanitized": result["sanitized"],
    }


@app.route("/", methods=["GET", "POST"])
def home():
    ctx = {
        "findings": [], "filename": "", "score": 0, "risk": "SAFE", "risk_class": "SAFE",
        "status": "AWAITING SCAN", "scanned_at": "", "error": "",
        "high_count": 0, "medium_count": 0, "low_count": 0,
        "summary": "Upload a document to assess its matching patterns.",
        "action_title": "Awaiting document", "action_text": "Choose a file to start the scan.",
        "verdict": "", "sanitized": "",
    }
    if request.method == "POST":
        upload = request.files.get("document")
        if not upload or not upload.filename:
            ctx.update(error="Select a file before scanning.", status="REJECTED")
        elif not upload.filename.lower().endswith(ALLOWED):
            ctx.update(error="Unsupported file type. Accepted: " + ", ".join(ALLOWED), status="REJECTED")
        else:
            name = secure_filename(upload.filename) or "document"
            try:
                result = analyze(name, upload.read())
            except Exception as exc:
                ctx.update(error=f"Could not read this file: {exc}", status="REJECTED")
            else:
                stored = False
                if result["verdict"] != "QUARANTINE":
                    APPROVED_STORE.append({"filename": name, "verdict": result["verdict"],
                                           "text": result["sanitized"]})
                    stored = True
                ctx = page_context(result, stored)
    return render_template("index.html", **ctx)


@app.errorhandler(413)
def too_large(_error):
    return "File too large. Maximum upload size is 5 MB.", 413


if __name__ == "__main__":
    app.run(debug=False, host="127.0.0.1", port=5000)

